try:
    from pptx import Presentation
except ImportError:
    import sys
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation

import sys

def extract_text(filename):
    prs = Presentation(filename)
    for i, slide in enumerate(prs.slides):
        print(f"--- Slide {i+1} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                print(shape.text)

if __name__ == '__main__':
    extract_text(sys.argv[1])
